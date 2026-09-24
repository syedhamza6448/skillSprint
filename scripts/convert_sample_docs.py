"""
Script to convert all 22 sample .txt documents into both .pdf and .docx formats.
"""

import os
import re
from pathlib import Path
import docx
from docx.shared import Pt, Inches, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

DOCS_DIR = Path(__file__).resolve().parents[1] / "sample_documents"

def txt_to_docx(txt_path: Path, docx_path: Path):
    doc = docx.Document()
    
    # Page setup
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    lines = txt_path.read_text(encoding="utf-8").splitlines()
    
    in_header = True
    body_lines = []
    
    for line in lines:
        stripped = line.strip()
        if not stripped:
            if body_lines:
                p = doc.add_paragraph(" ".join(body_lines))
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                body_lines = []
            continue

        # Skip decorative divider lines
        if re.match(r"^[=\-]{5,}$", stripped):
            if body_lines:
                p = doc.add_paragraph(" ".join(body_lines))
                p.paragraph_format.space_after = Pt(6)
                body_lines = []
            continue

        # Major section
        sec_m = re.match(r"^SECTION\s+\d+:\s*(.*)", stripped, re.IGNORECASE)
        if sec_m:
            if body_lines:
                p = doc.add_paragraph(" ".join(body_lines))
                p.paragraph_format.space_after = Pt(6)
                body_lines = []
            h = doc.add_heading(stripped, level=1)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)
            continue

        # Subsection e.g. "3.1 Working Hours"
        sub_m = re.match(r"^(\d+\.\d+)\s+(.*)", stripped)
        if sub_m:
            if body_lines:
                p = doc.add_paragraph(" ".join(body_lines))
                p.paragraph_format.space_after = Pt(6)
                body_lines = []
            h = doc.add_heading(stripped, level=2)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(4)
            continue

        # Header block check (Title, Document ID, Version)
        if in_header and (
            "NEXACORE SOLUTIONS" in stripped
            or "Document ID:" in stripped
            or "Version:" in stripped
            or "Effective Date:" in stripped
            or "Supersedes:" in stripped
            or stripped.isupper() and len(stripped) < 50
        ):
            p = doc.add_paragraph(stripped)
            p.paragraph_format.space_after = Pt(2)
            p.runs[0].bold = True
            continue
        else:
            in_header = False

        body_lines.append(stripped)

    if body_lines:
        p = doc.add_paragraph(" ".join(body_lines))
        p.paragraph_format.space_after = Pt(6)

    doc.save(docx_path)


def txt_to_pdf(txt_path: Path, pdf_path: Path):
    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#1A365D"),
        spaceAfter=6,
    )
    
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#4A5568"),
        spaceAfter=2,
    )

    h1_style = ParagraphStyle(
        'SectionH1',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#2B6CB0"),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True,
    )

    h2_style = ParagraphStyle(
        'SectionH2',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#2D3748"),
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True,
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#1A202C"),
        spaceAfter=6,
    )

    story = []
    lines = txt_path.read_text(encoding="utf-8").splitlines()
    in_header = True
    body_lines = []

    def flush_body():
        nonlocal body_lines
        if body_lines:
            text = " ".join(body_lines)
            # escape xml special chars for reportlab
            text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(text, body_style))
            body_lines = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush_body()
            continue

        if re.match(r"^[=\-]{5,}$", stripped):
            flush_body()
            continue

        sec_m = re.match(r"^SECTION\s+\d+:\s*(.*)", stripped, re.IGNORECASE)
        if sec_m:
            flush_body()
            text = stripped.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(text, h1_style))
            continue

        sub_m = re.match(r"^(\d+\.\d+)\s+(.*)", stripped)
        if sub_m:
            flush_body()
            text = stripped.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(text, h2_style))
            continue

        if in_header and (
            "NEXACORE SOLUTIONS" in stripped
            or "Document ID:" in stripped
            or "Version:" in stripped
            or "Effective Date:" in stripped
            or "Supersedes:" in stripped
            or stripped.isupper() and len(stripped) < 50
        ):
            text = stripped.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if "NEXACORE SOLUTIONS" in stripped or (stripped.isupper() and "DOCUMENT" in stripped or "HANDBOOK" in stripped or "POLICY" in stripped or "SOP" in stripped or "MANUAL" in stripped or "GUIDELINES" in stripped or "FAQ" in stripped):
                story.append(Paragraph(text, title_style))
            else:
                story.append(Paragraph(text, meta_style))
            continue
        else:
            in_header = False

        body_lines.append(stripped)

    flush_body()
    doc.build(story)


def main():
    txt_files = sorted(list(DOCS_DIR.glob("*.txt")))
    print(f"Found {len(txt_files)} text files in {DOCS_DIR}")
    for f in txt_files:
        stem = f.stem
        docx_path = DOCS_DIR / f"{stem}.docx"
        pdf_path = DOCS_DIR / f"{stem}.pdf"
        
        print(f"Converting {f.name} -> .docx and .pdf...")
        txt_to_docx(f, docx_path)
        txt_to_pdf(f, pdf_path)
    print("All conversions complete!")

if __name__ == "__main__":
    main()
