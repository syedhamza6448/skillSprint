import os
import glob
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

def convert_to_docx(md_path, docx_path):
    doc = Document()
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        else:
            doc.add_paragraph(line)
    
    doc.save(docx_path)

def convert_to_pdf(md_path, pdf_path):
    doc = SimpleDocTemplate(pdf_path, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip()
        if not line:
            story.append(Spacer(1, 12))
            continue
        if line.startswith('### '):
            story.append(Paragraph(line[4:], styles['Heading3']))
        elif line.startswith('## '):
            story.append(Paragraph(line[3:], styles['Heading2']))
        elif line.startswith('# '):
            story.append(Paragraph(line[2:], styles['Heading1']))
        else:
            story.append(Paragraph(line, styles['Normal']))
            
    doc.build(story)

def main():
    drafts_dir = 'data/company_docs_drafts'
    out_dir = 'data/company_docs'
    os.makedirs(out_dir, exist_ok=True)
    
    md_files = sorted(glob.glob(os.path.join(drafts_dir, '*.md')))
    
    for i, md_file in enumerate(md_files):
        filename = os.path.basename(md_file)
        name, _ = os.path.splitext(filename)
        
        if i < 6:
            # First half to PDF
            out_path = os.path.join(out_dir, name + '.pdf')
            convert_to_pdf(md_file, out_path)
            print(f"Converted to PDF: {out_path}")
        else:
            # Second half to DOCX
            out_path = os.path.join(out_dir, name + '.docx')
            convert_to_docx(md_file, out_path)
            print(f"Converted to DOCX: {out_path}")

if __name__ == '__main__':
    main()
